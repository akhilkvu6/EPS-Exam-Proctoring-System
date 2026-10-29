from django.test import TestCase, Client
from django.urls import reverse
from accounts.models import (
    User, TeacherProfile, StudentProfile, Department,
    ClassGroup, Notification, ApprovalRecord, AuditLog
)


class AccountsAuthenticationAndApprovalTests(TestCase):
    def setUp(self):
        self.client = Client()

        # Create department and class group
        self.dept = Department.objects.create(name='Computer Science', code='CS')
        self.cls_grp = ClassGroup.objects.create(department=self.dept, name='BCA', semester='3')

        # Admin user
        self.admin_user = User.objects.create_superuser(
            username='admin_eps',
            email='admin@eps.edu',
            password='AdminPassword123!',
            role='admin'
        )

        # Approved teacher
        self.teacher_user = User.objects.create_user(
            username='prof_sharma',
            email='sharma@eps.edu',
            password='TeacherPassword123!',
            role='teacher',
            first_name='Prof. Sharma'
        )
        self.teacher_profile = TeacherProfile.objects.create(
            user=self.teacher_user,
            employee_id='EMP101',
            department='Computer Science',
            qualification='Ph.D',
            verification_status='approved'
        )
        self.teacher_profile.assigned_classes.add(self.cls_grp)

        # Approved student
        self.student_user = User.objects.create_user(
            username='student_rahul',
            email='rahul@eps.edu',
            password='StudentPassword123!',
            role='student',
            first_name='Rahul Kumar'
        )
        self.student_profile = StudentProfile.objects.create(
            user=self.student_user,
            full_name='Rahul Kumar',
            roll_number='BCA202601',
            department='Computer Science',
            semester='3',
            class_group=self.cls_grp,
            verification_status='approved'
        )

    def test_student_registration_workflow(self):
        """Student registration creates account in Pending status and notifies department teachers."""
        initial_notif_count = Notification.objects.filter(recipient=self.teacher_user).count()

        resp = self.client.post(reverse('student_register'), {
            'username': 'new_student_priya',
            'full_name': 'Priya Patel',
            'email': 'priya@eps.edu',
            'roll_number': 'BCA202602',
            'department': 'Computer Science',
            'semester': '3',
            'password1': 'SecurePass123!',
            'password2': 'SecurePass123!',
        })
        self.assertEqual(resp.status_code, 302)
        self.assertRedirects(resp, reverse('student_registration_success'))

        new_student = User.objects.get(username='new_student_priya')
        self.assertEqual(new_student.role, 'student')
        self.assertEqual(new_student.student_profile.verification_status, 'pending')

        # Verify teacher notification was created
        new_notif_count = Notification.objects.filter(recipient=self.teacher_user).count()
        self.assertEqual(new_notif_count, initial_notif_count + 1)

    def test_teacher_approves_student(self):
        """Teacher can review and approve a pending student."""
        # Create pending student
        pending_user = User.objects.create_user(
            username='pending_student_amit',
            password='Pass12345!',
            role='student'
        )
        pending_profile = StudentProfile.objects.create(
            user=pending_user,
            full_name='Amit Singh',
            roll_number='BCA202699',
            department='Computer Science',
            semester='3',
            verification_status='pending'
        )

        self.client.force_login(self.teacher_user)
        resp = self.client.post(reverse('review_student', args=[pending_profile.id]), {
            'action': 'approve',
            'reason': 'Verified enrollment records.'
        })
        self.assertEqual(resp.status_code, 302)

        pending_profile.refresh_from_db()
        self.assertEqual(pending_profile.verification_status, 'approved')
        self.assertEqual(pending_profile.reviewed_by, self.teacher_user)

        # Check ApprovalRecord created
        self.assertTrue(ApprovalRecord.objects.filter(applicant=pending_user, decision='approved').exists())

        # Check student received in-app notification
        self.assertTrue(Notification.objects.filter(recipient=pending_user, notification_type='approval').exists())

    def test_teacher_rejects_student(self):
        """Teacher can reject a student application with reason."""
        pending_user = User.objects.create_user(
            username='fake_student',
            password='Pass12345!',
            role='student'
        )
        pending_profile = StudentProfile.objects.create(
            user=pending_user,
            full_name='Fake User',
            roll_number='BCA999999',
            department='Computer Science',
            semester='3',
            verification_status='pending'
        )

        self.client.force_login(self.teacher_user)
        resp = self.client.post(reverse('review_student', args=[pending_profile.id]), {
            'action': 'reject',
            'reason': 'Roll number not in university ledger.'
        })
        self.assertEqual(resp.status_code, 302)

        pending_profile.refresh_from_db()
        self.assertEqual(pending_profile.verification_status, 'rejected')
        self.assertIn('ledger', pending_profile.rejection_reason)

    def test_teacher_registration_and_admin_approval(self):
        """Teacher registration is pending until administrator approves."""
        resp = self.client.post(reverse('teacher_register'), {
            'username': 'prof_anita',
            'full_name': 'Prof. Anita Roy',
            'email': 'anita@eps.edu',
            'employee_id': 'EMP202',
            'department': 'Computer Science',
            'qualification': 'M.Tech',
            'institution': 'State Institute',
            'password1': 'TeacherPass123!',
            'password2': 'TeacherPass123!',
        })
        self.assertEqual(resp.status_code, 302)

        new_teacher = User.objects.get(username='prof_anita')
        self.assertEqual(new_teacher.teacher_profile.verification_status, 'pending')

        # Admin notified
        self.assertTrue(Notification.objects.filter(recipient=self.admin_user, notification_type='registration').exists())

        # Admin approves teacher
        self.client.force_login(self.admin_user)
        resp = self.client.post(reverse('review_teacher', args=[new_teacher.teacher_profile.id]), {
            'action': 'approve',
            'reason': 'Verified credentials.'
        })
        self.assertEqual(resp.status_code, 302)

        new_teacher.teacher_profile.refresh_from_db()
        self.assertEqual(new_teacher.teacher_profile.verification_status, 'approved')

    def test_common_login_role_redirection(self):
        """Common login correctly directs Student, Teacher, and Admin."""
        # 1. Student login
        resp = self.client.post(reverse('common_login'), {
            'username': 'student_rahul',
            'password': 'StudentPassword123!'
        })
        self.assertRedirects(resp, reverse('student_dashboard'))
        self.client.logout()

        # 2. Teacher login
        resp = self.client.post(reverse('common_login'), {
            'username': 'prof_sharma',
            'password': 'TeacherPassword123!'
        })
        self.assertRedirects(resp, reverse('teacher_dashboard'))
        self.client.logout()

        # 3. Admin login
        resp = self.client.post(reverse('common_login'), {
            'username': 'admin_eps',
            'password': 'AdminPassword123!'
        })
        self.assertRedirects(resp, reverse('admin_dashboard'))
        self.client.logout()

    def test_unapproved_student_blocked_from_portal(self):
        """Pending student cannot access student_dashboard; redirected to account_status."""
        pending_user = User.objects.create_user(
            username='student_waiting',
            password='Pass12345!',
            role='student'
        )
        StudentProfile.objects.create(
            user=pending_user,
            roll_number='BCA5555',
            department='Computer Science',
            semester='3',
            verification_status='pending'
        )

        self.client.force_login(pending_user)
        resp = self.client.get(reverse('student_dashboard'))
        self.assertRedirects(resp, reverse('account_status'))

    def test_unauthorized_dashboard_access_protection(self):
        """Students cannot access teacher_dashboard or admin_dashboard."""
        self.client.force_login(self.student_user)

        # Student accessing teacher dashboard is redirected to student dashboard
        resp = self.client.get(reverse('teacher_dashboard'))
        self.assertRedirects(resp, reverse('student_dashboard'))

        # Student accessing admin dashboard is redirected to student dashboard
        resp = self.client.get(reverse('admin_dashboard'))
        self.assertRedirects(resp, reverse('student_dashboard'))

        # Anonymous user is redirected to common login
        self.client.logout()
        resp = self.client.get(reverse('teacher_dashboard'))
        self.assertRedirects(resp, reverse('common_login'))

    def test_in_app_password_change(self):
        """Users can securely change password in-app."""
        self.client.force_login(self.student_user)
        resp = self.client.post(reverse('change_password'), {
            'current_password': 'StudentPassword123!',
            'new_password': 'BrandNewPassword123!',
            'confirm_password': 'BrandNewPassword123!',
        })
        self.assertEqual(resp.status_code, 302)

        # Check authentication with new password
        self.client.logout()
        resp = self.client.post(reverse('common_login'), {
            'username': 'student_rahul',
            'password': 'BrandNewPassword123!'
        })
        self.assertRedirects(resp, reverse('student_dashboard'))

    def test_notifications_marked_read(self):
        """In-app notifications mark read works."""
        n = Notification.objects.create(
            recipient=self.student_user,
            title='Test Alert',
            message='Testing notification delivery',
            notification_type='system'
        )
        self.assertFalse(n.is_read)

        self.client.force_login(self.student_user)
        resp = self.client.get(reverse('mark_notification_read', args=[n.id]))
        n.refresh_from_db()
        self.assertTrue(n.is_read)


class ExcelBulkStudentImportTests(TestCase):
    def setUp(self):
        self.client = Client()

        self.dept = Department.objects.create(name='Computer Science', code='CS')

        # Teacher
        self.teacher_user = User.objects.create_user(
            username='prof_kapoor',
            password='Password123!',
            role='teacher',
            first_name='Prof. Kapoor'
        )
        self.teacher_profile = TeacherProfile.objects.create(
            user=self.teacher_user,
            employee_id='EMP999',
            department='Computer Science',
            verification_status='approved'
        )

        # Existing student
        self.existing_student = User.objects.create_user(
            username='existing_student',
            password='Password123!',
            role='student'
        )
        StudentProfile.objects.create(
            user=self.existing_student,
            full_name='Existing Student',
            roll_number='CS2026_EXIST',
            department='Computer Science',
            semester='3',
            verification_status='approved'
        )

    def _create_excel_file(self, rows):
        import io
        import openpyxl
        from django.core.files.uploadedfile import SimpleUploadedFile

        headers = ["Register Number", "Full Name", "Username", "Initial Password", "Department", "Semester", "Contact Number"]
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(headers)
        for r in rows:
            ws.append(r)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return SimpleUploadedFile("students.xlsx", buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def test_download_template_access_and_format(self):
        """Only staff/faculty can download the template; verifies valid Excel sheet."""
        import openpyxl
        import io

        # Anonymous cannot download
        resp = self.client.get(reverse('download_student_template'))
        self.assertRedirects(resp, reverse('common_login'))

        # Teacher can download
        self.client.force_login(self.teacher_user)
        resp = self.client.get(reverse('download_student_template'))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        # Verify openpyxl can read it
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        headers = [c.value for c in ws[1]]
        self.assertIn("Register Number", headers)
        self.assertIn("Username", headers)

    def test_bulk_student_import_success(self):
        """Teacher successfully bulk imports valid students in pre-approved status."""
        self.client.force_login(self.teacher_user)

        rows = [
            ["BCA101", "Aakash Singh", "aakash_s", "StudentPass1!", "Computer Science", "3", "9998887771"],
            ["BCA102", "Pooja Hegde", "pooja_h", "StudentPass2!", "Computer Science", "3", "9998887772"],
        ]
        excel_file = self._create_excel_file(rows)

        resp = self.client.post(reverse('import_students'), {
            'excel_file': excel_file,
            'initial_status': 'approved',
            'import_mode': 'skip_errors',
            'default_department': 'Computer Science',
        })
        self.assertEqual(resp.status_code, 200)

        # Verify students created
        self.assertTrue(User.objects.filter(username='aakash_s').exists())
        self.assertTrue(User.objects.filter(username='pooja_h').exists())

        p = StudentProfile.objects.get(roll_number='BCA101')
        self.assertEqual(p.full_name, 'Aakash Singh')
        self.assertEqual(p.verification_status, 'approved')
        self.assertEqual(p.reviewed_by, self.teacher_user)

        # Check in-app notification delivered to imported student (strictly email-free)
        aakash_user = User.objects.get(username='aakash_s')
        self.assertTrue(aakash_user.notifications.filter(title__icontains='Welcome').exists())

        # Check in-app notification delivered to teacher
        self.assertTrue(self.teacher_user.notifications.filter(title__icontains='Bulk Student Import').exists())

    def test_bulk_student_import_duplicate_and_error_handling(self):
        """Valid rows import while duplicates and errors are reported cleanly."""
        self.client.force_login(self.teacher_user)

        rows = [
            # Valid row
            ["BCA201", "Divya Nair", "divya_n", "DivyaPass123!", "Computer Science", "3", "9991112223"],
            # Duplicate username against DB
            ["BCA202", "Duplicate User", "existing_student", "TestPass123!", "Computer Science", "3", "9991112224"],
            # Duplicate roll number against DB
            ["CS2026_EXIST", "Duplicate Roll", "divya_dup", "TestPass123!", "Computer Science", "3", "9991112225"],
            # Password too short (<6 chars)
            ["BCA204", "Short Pass", "short_p", "123", "Computer Science", "3", "9991112226"],
        ]
        excel_file = self._create_excel_file(rows)

        resp = self.client.post(reverse('import_students'), {
            'excel_file': excel_file,
            'initial_status': 'approved',
            'import_mode': 'skip_errors',
            'default_department': 'Computer Science',
        })
        self.assertEqual(resp.status_code, 200)

        # Only divya_n was created
        self.assertTrue(User.objects.filter(username='divya_n').exists())
        self.assertFalse(User.objects.filter(username='short_p').exists())

        summary = resp.context['import_summary']
        self.assertEqual(summary['imported_count'], 1)
        self.assertEqual(summary['error_count'], 3)

    def test_bulk_student_import_atomic_rollback(self):
        """In atomic mode, any error rolls back the entire batch."""
        self.client.force_login(self.teacher_user)

        rows = [
            # Valid student
            ["BCA301", "Karan Johar", "karan_j", "KaranPass123!", "Computer Science", "3", "9993334441"],
            # Errored student (password too short)
            ["BCA302", "Bad Row", "bad_row", "12", "Computer Science", "3", "9993334442"],
        ]
        excel_file = self._create_excel_file(rows)

        resp = self.client.post(reverse('import_students'), {
            'excel_file': excel_file,
            'initial_status': 'approved',
            'import_mode': 'atomic',
            'default_department': 'Computer Science',
        })
        self.assertEqual(resp.status_code, 200)

        # Atomic rollback: karan_j must NOT be created
        self.assertFalse(User.objects.filter(username='karan_j').exists())



class CentralRegisterAndAdminPortalTests(TestCase):
    def setUp(self):
        self.client = Client()
        self.dept = Department.objects.create(name='Computer Science', code='CS')
        self.cls_grp = ClassGroup.objects.create(department=self.dept, name='BCA-A', semester='3')

        # Admin user
        self.admin = User.objects.create_superuser(
            username='system_admin',
            email='admin@eps.edu',
            password='AdminPassword123!',
            role='admin'
        )

        # Faculty user
        self.teacher_user = User.objects.create_user(
            username='prof_gupta',
            email='gupta@eps.edu',
            password='GuptaPassword123!',
            role='teacher',
            first_name='Prof. Gupta'
        )
        self.teacher_profile = TeacherProfile.objects.create(
            user=self.teacher_user,
            employee_id='EMP9901',
            department='Computer Science',
            qualification='M.Tech',
            verification_status='approved'
        )

        # Student user
        self.student_user = User.objects.create_user(
            username='student_rohit',
            email='rohit@eps.edu',
            password='RohitPassword123!',
            role='student',
            first_name='Rohit Verma'
        )
        self.student_profile = StudentProfile.objects.create(
            user=self.student_user,
            full_name='Rohit Verma',
            roll_number='BCA202688',
            department='Computer Science',
            semester='3',
            class_group=self.cls_grp,
            verification_status='approved'
        )

    def test_central_register_get_renders_role_options(self):
        """GET /accounts/register/ returns 200 and presents role options."""
        resp = self.client.get(reverse('central_register'))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Which type of account do you need?')
        self.assertContains(resp, 'Student')
        self.assertContains(resp, 'Teacher')

    def test_central_register_student_form_submission(self):
        """Submitting student registration via central register flow creates pending student."""
        resp = self.client.post(reverse('central_register'), {
            'role': 'student',
            'full_name': 'Meera Sen',
            'username': 'meera_sen',
            'email': 'meera@eps.edu',
            'password1': 'MeeraPass123!',
            'password2': 'MeeraPass123!',
            'roll_number': 'BCA202677',
            'department': 'Computer Science',
            'semester': '3',
            'class_group': self.cls_grp.id,
        })
        self.assertEqual(resp.status_code, 302)
        self.assertRedirects(resp, reverse('registration_success'))

        user = User.objects.get(username='meera_sen')
        self.assertEqual(user.role, 'student')
        self.assertEqual(user.student_profile.roll_number, 'BCA202677')
        self.assertEqual(user.student_profile.verification_status, 'pending')

    def test_central_register_teacher_form_submission(self):
        """Submitting teacher registration via central register flow creates pending teacher."""
        resp = self.client.post(reverse('central_register'), {
            'role': 'teacher',
            'full_name': 'Dr. Raman',
            'username': 'dr_raman',
            'email': 'raman@eps.edu',
            'password1': 'RamanPass123!',
            'password2': 'RamanPass123!',
            'employee_id': 'EMP7766',
            'department': 'Computer Science',
            'qualification': 'Ph.D in AI',
            'institution': 'EPS Engineering Institute',
        })
        self.assertEqual(resp.status_code, 302)
        self.assertRedirects(resp, reverse('registration_success'))

        user = User.objects.get(username='dr_raman')
        self.assertEqual(user.role, 'teacher')
        self.assertEqual(user.teacher_profile.employee_id, 'EMP7766')
        self.assertEqual(user.teacher_profile.verification_status, 'pending')

    def test_admin_portal_access_control(self):
        """Non-admin users cannot access admin endpoints; admin can access all."""
        # Unauthenticated access redirects to login
        resp = self.client.get(reverse('admin_dashboard'))
        self.assertEqual(resp.status_code, 302)

        # Student access redirects to student dashboard or 403
        self.client.force_login(self.student_user)
        resp = self.client.get(reverse('admin_dashboard'))
        self.assertIn(resp.status_code, [302, 403])

        # Teacher access redirects or 403
        self.client.force_login(self.teacher_user)
        resp = self.client.get(reverse('admin_dashboard'))
        self.assertIn(resp.status_code, [302, 403])

        # Admin access succeeds with HTTP 200 across all admin routes
        self.client.force_login(self.admin)
        admin_urls = [
            'admin_dashboard',
            'admin_analytics',
            'admin_audit_logs',
            'admin_students',
            'admin_teachers',
            'admin_student_approvals',
            'admin_teacher_approvals',
            'admin_suspended_accounts',
            'admin_approval_history',
            'admin_departments',
            'admin_classes',
            'admin_policies',
            'admin_backup',
            'admin_security',
            'admin_all_exams',
            'admin_exam_attempts',
            'admin_proctoring_overview',
            'admin_proctoring_incidents',
            'admin_results_overview',
            'admin_results_export',
        ]
        for url_name in admin_urls:
            r = self.client.get(reverse(url_name))
            self.assertEqual(r.status_code, 200, f"Route {url_name} failed to return 200 for admin: {r.status_code}")

    def test_admin_student_directory_search_and_filter(self):
        """Admin can search and filter students in directory."""
        self.client.force_login(self.admin)
        resp = self.client.get(reverse('admin_students'), {'q': 'Rohit'})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Rohit Verma')

        # Filter by department
        resp = self.client.get(reverse('admin_students'), {'department': 'Computer Science'})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Rohit Verma')

        # Filter by status
        resp = self.client.get(reverse('admin_students'), {'status': 'approved'})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Rohit Verma')

    def test_admin_bulk_suspend_and_activate_students(self):
        """Admin can bulk suspend and activate student accounts."""
        self.client.force_login(self.admin)

        # Bulk suspend
        resp = self.client.post(reverse('admin_students'), {
            'bulk_action': 'suspend',
            'selected_students': [self.student_profile.id],
        })
        self.assertEqual(resp.status_code, 302)
        self.student_profile.refresh_from_db()
        self.assertEqual(self.student_profile.verification_status, 'suspended')

        # Bulk activate
        resp = self.client.post(reverse('admin_students'), {
            'bulk_action': 'activate',
            'selected_students': [self.student_profile.id],
        })
        self.assertEqual(resp.status_code, 302)
        self.student_profile.refresh_from_db()
        self.assertEqual(self.student_profile.verification_status, 'approved')

    def test_admin_review_student_approve_and_reject(self):
        """Admin can approve and reject pending student applications from the oversight queue."""
        self.client.force_login(self.admin)

        # Create pending student
        p_user = User.objects.create_user(username='pending_stu_1', password='Pass12345!', role='student')
        p_prof = StudentProfile.objects.create(
            user=p_user, roll_number='BCA991', department='Computer Science', semester='3', verification_status='pending'
        )

        # Approve student
        resp = self.client.post(reverse('admin_review_student', args=[p_prof.id]), {
            'action': 'approve',
            'reason': 'Approved by administrator test',
        })
        self.assertEqual(resp.status_code, 302)
        p_prof.refresh_from_db()
        self.assertEqual(p_prof.verification_status, 'approved')
        self.assertTrue(ApprovalRecord.objects.filter(applicant=p_user, decision='approved').exists())

        # Reject student
        resp = self.client.post(reverse('admin_review_student', args=[p_prof.id]), {
            'action': 'reject',
            'reason': 'Roll number mismatch',
        })
        self.assertEqual(resp.status_code, 302)
        p_prof.refresh_from_db()
        self.assertEqual(p_prof.verification_status, 'rejected')
        self.assertTrue(ApprovalRecord.objects.filter(applicant=p_user, decision='rejected').exists())

    def test_admin_review_teacher_approve_and_reject(self):
        """Admin can approve and reject faculty applications."""
        self.client.force_login(self.admin)

        p_t_user = User.objects.create_user(username='pending_tea_1', password='Pass12345!', role='teacher')
        p_t_prof = TeacherProfile.objects.create(
            user=p_t_user, employee_id='EMP9988', department='Computer Science', verification_status='pending'
        )

        # Approve
        resp = self.client.post(reverse('review_teacher', args=[p_t_prof.id]), {
            'action': 'approve',
            'reason': 'Verified credential',
        })
        self.assertEqual(resp.status_code, 302)
        p_t_prof.refresh_from_db()
        self.assertEqual(p_t_prof.verification_status, 'approved')

        # Reject
        resp = self.client.post(reverse('review_teacher', args=[p_t_prof.id]), {
            'action': 'reject',
            'reason': 'Invalid employee ID',
        })
        self.assertEqual(resp.status_code, 302)
        p_t_prof.refresh_from_db()
        self.assertEqual(p_t_prof.verification_status, 'rejected')

    def test_admin_export_reports_csv(self):
        """Admin can download CSV exports for students, teachers, results, incidents, and audit logs."""
        self.client.force_login(self.admin)

        reports = ['students', 'teachers', 'results', 'incidents', 'audit']
        for r_type in reports:
            resp = self.client.get(reverse('admin_results_export'), {'report': r_type})
            self.assertEqual(resp.status_code, 200)
            self.assertEqual(resp['Content-Type'], 'text/csv')
            self.assertIn('attachment;', resp['Content-Disposition'])

    def test_teacher_student_approvals_queue_view(self):
        """Teacher can access their departmental student approval queue."""
        self.client.force_login(self.teacher_user)
        resp = self.client.get(reverse('teacher_student_approvals'))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Candidate Approval Queue')
        self.assertContains(resp, 'Computer Science Department')


class InstitutionalPoliciesTests(TestCase):
    def setUp(self):
        self.dept = Department.objects.create(name='Computer Science', code='CS')
        self.admin = User.objects.create_superuser(
            username='admin_policy_test',
            password='AdminPassword123!',
            email='admin@eps.edu',
            role='admin'
        )
        self.student = User.objects.create_user(
            username='student_policy_test',
            password='StudentPassword123!',
            email='student@eps.edu',
            role='student'
        )
        StudentProfile.objects.create(
            user=self.student,
            full_name='Test Student',
            roll_number='BCA202699',
            department='Computer Science',
            semester='3',
            verification_status='approved'
        )
        self.teacher = User.objects.create_user(
            username='teacher_policy_test',
            password='TeacherPassword123!',
            email='teacher@eps.edu',
            role='teacher'
        )
        TeacherProfile.objects.create(
            user=self.teacher,
            employee_id='EMP999',
            department='Computer Science',
            verification_status='approved'
        )

    def test_unauthorized_access_blocked(self):
        # Anonymous
        resp = self.client.get(reverse('admin_policies'))
        self.assertEqual(resp.status_code, 302)
        
        # Student
        self.client.force_login(self.student)
        resp = self.client.get(reverse('admin_policies'))
        self.assertEqual(resp.status_code, 302)
        
        # Teacher
        self.client.force_login(self.teacher)
        resp = self.client.get(reverse('admin_policies'))
        self.assertEqual(resp.status_code, 302)

    def test_admin_policies_get_renders_correctly(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse('admin_policies'))
        self.assertEqual(resp.status_code, 200)
        self.assertTemplateUsed(resp, 'accounts/admin_policies.html')
        self.assertContains(resp, 'Examination and Proctoring Policies')
        self.assertContains(resp, 'Active Policy Hash:')
        self.assertContains(resp, 'Global Enforcement')
        self.assertIn('policies', resp.context)
        self.assertIn('metadata', resp.context)

    def test_admin_policies_post_saves_and_persists(self):
        self.client.force_login(self.admin)
        post_data = {
            'late_entry_mins': 25,
            'late_entry_enabled': 'on',
            'early_entry_enabled': 'on',
            'attempt_allowance': 'dual',
            'server_submission_enabled': 'on',
            'timing_accommodation': '1.5',
            'video_ingestion': 'single',
            'audio_monitor': 'continuous',
            'os_lockdown_enabled': 'on',
            'focus_trap_enabled': 'on',
            'clipboard_neutralization_enabled': 'on',
            'item_randomization_enabled': 'on',
            'distractor_shuffling_enabled': 'on',
            'negative_marking_scheme': 'strict',
            'score_embargo_enabled': 'on',
            'telemetry_retention_days': '180',
            'audit_retention_years': '10y',
        }
        resp = self.client.post(reverse('admin_policies'), post_data)
        self.assertEqual(resp.status_code, 302)
        self.assertRedirects(resp, reverse('admin_policies'))
        
        # Verify persistence on subsequent GET
        get_resp = self.client.get(reverse('admin_policies'))
        self.assertEqual(get_resp.status_code, 200)
        saved_policies = get_resp.context['policies']
        self.assertEqual(saved_policies['late_entry_mins'], 25)
        self.assertEqual(saved_policies['attempt_allowance'], 'dual')
        self.assertEqual(saved_policies['timing_accommodation'], '1.5')
        self.assertEqual(saved_policies['video_ingestion'], 'single')
        self.assertEqual(saved_policies['audio_monitor'], 'continuous')
        self.assertEqual(saved_policies['negative_marking_scheme'], 'strict')
        self.assertEqual(saved_policies['telemetry_retention_days'], '180')
        self.assertEqual(saved_policies['audit_retention_years'], '10y')
        
        # Verify AuditLog created
        audit_entry = AuditLog.objects.filter(action='POLICY_UPDATE').last()
        self.assertIsNotNone(audit_entry)
        self.assertEqual(audit_entry.actor, self.admin)

    def test_admin_policies_export_manifest_json(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse('admin_policies_export'))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/json')
        data = resp.json()
        self.assertEqual(data['manifest_version'], '4.8')
        self.assertIn('policy_sha256_hash', data)
        self.assertIn('institutional_policies', data)
        self.assertTrue(data['stratum_ntp_sync'])


class ProfileAndTeacherRosterTests(TestCase):
    def setUp(self):
        self.client = Client()
        self.dept = Department.objects.create(name='Computer Science', code='CS')
        self.cls_grp = ClassGroup.objects.create(department=self.dept, name='BCA 3A', semester='3')
        self.cls_grp_b = ClassGroup.objects.create(department=self.dept, name='BCA 3B', semester='3')

        self.student = User.objects.create_user(
            username='student_test',
            email='student@eps.edu',
            password='TestPass123!',
            role='student',
            first_name='Original Student'
        )
        self.student_prof = StudentProfile.objects.create(
            user=self.student,
            full_name='Original Student',
            roll_number='ROLL100',
            department='Computer Science',
            semester='3',
            class_group=self.cls_grp,
            verification_status='approved'
        )

        self.teacher = User.objects.create_user(
            username='teacher_test',
            email='teacher@eps.edu',
            password='TestPass123!',
            role='teacher',
            first_name='Dr. Teacher'
        )
        self.teacher_prof = TeacherProfile.objects.create(
            user=self.teacher,
            employee_id='EMP200',
            department='Computer Science',
            qualification='M.Tech',
            institution='EPS University',
            verification_status='approved'
        )

        self.admin = User.objects.create_superuser(
            username='admin_test',
            email='admin@eps.edu',
            password='TestPass123!',
            role='admin'
        )

    def test_anonymous_profile_redirects_to_login(self):
        resp = self.client.get(reverse('user_profile'))
        self.assertEqual(resp.status_code, 302)
        self.assertIn('/accounts/login/', resp.url)

    def test_student_can_view_and_edit_profile(self):
        self.client.force_login(self.student)
        resp = self.client.get(reverse('user_profile'))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'ROLL100')
        self.assertContains(resp, 'Computer Science')

        # Edit student profile
        post_resp = self.client.post(reverse('user_profile'), {
            'full_name': 'Updated Student Name',
            'email': 'updated_student@eps.edu',
            'semester': '4',
        })
        self.assertEqual(post_resp.status_code, 302)
        self.student.refresh_from_db()
        self.student_prof.refresh_from_db()
        self.assertEqual(self.student.first_name, 'Updated Student Name')
        self.assertEqual(self.student.email, 'updated_student@eps.edu')
        self.assertEqual(self.student_prof.semester, '4')

    def test_teacher_can_view_and_edit_profile(self):
        self.client.force_login(self.teacher)
        resp = self.client.get(reverse('user_profile'))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'EMP200')

        post_resp = self.client.post(reverse('user_profile'), {
            'full_name': 'Prof. Senior Teacher',
            'email': 'senior_teacher@eps.edu',
            'department': 'Computer Science',
            'qualification': 'Ph.D in AI',
            'institution': 'EPS Campus North',
        })
        self.assertEqual(post_resp.status_code, 302)
        self.teacher.refresh_from_db()
        self.teacher_prof.refresh_from_db()
        self.assertEqual(self.teacher.first_name, 'Prof. Senior Teacher')
        self.assertEqual(self.teacher_prof.qualification, 'Ph.D in AI')
        self.assertEqual(self.teacher_prof.institution, 'EPS Campus North')

    def test_teacher_students_roster_and_class_assignment(self):
        self.client.force_login(self.teacher)
        resp = self.client.get(reverse('teacher_students'))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'ROLL100')

        # Filter by class
        resp_filter = self.client.get(reverse('teacher_students') + f'?class_filter={self.cls_grp.id}')
        self.assertEqual(resp_filter.status_code, 200)
        self.assertContains(resp_filter, 'ROLL100')

        # Reassign student to Class Group B
        post_resp = self.client.post(reverse('teacher_students'), {
            'assign_class': '1',
            'student_id': self.student_prof.id,
            'class_id': self.cls_grp_b.id,
        })
        self.assertEqual(post_resp.status_code, 302)
        self.student_prof.refresh_from_db()
        self.assertEqual(self.student_prof.class_group, self.cls_grp_b)

        # AuditLog logged
        self.assertTrue(AuditLog.objects.filter(action='STUDENT_CLASS_ASSIGNED').exists())


class SecurityRemediationStage1Tests(TestCase):
    """
    Regression tests verifying Stage 1 Security Remediation:
    - SEC-03: Teacher cannot change their own department via profile
    - SEC-02: Strict department scoping for student approvals
    - SEC-01 & SEC-14: Production configuration hardening (Secret Key, Allowed Hosts, Cookies, HSTS)
    """

    def setUp(self):
        self.client = Client()

        # Departments
        self.dept_cs = Department.objects.create(name='Computer Science', code='CS')
        self.dept_ce = Department.objects.create(name='Civil Engineering', code='CE')

        # Class groups
        self.cls_cs = ClassGroup.objects.create(department=self.dept_cs, name='CS-1', semester='1')
        self.cls_ce = ClassGroup.objects.create(department=self.dept_ce, name='CE-1', semester='1')

        # Admin user
        self.admin_user = User.objects.create_superuser(
            username='admin_stage1',
            email='admin_stage1@eps.edu',
            password='AdminPassword123!',
            role='admin'
        )

        # Teacher - Computer Science
        self.teacher_cs = User.objects.create_user(
            username='teacher_cs_user',
            email='teacher_cs@eps.edu',
            password='TeacherPass123!',
            role='teacher',
            first_name='Prof. CS'
        )
        self.teacher_cs_profile = TeacherProfile.objects.create(
            user=self.teacher_cs,
            employee_id='EMP_CS_01',
            department='Computer Science',
            qualification='M.Tech CS',
            institution='Main Campus',
            verification_status='approved'
        )

        # Teacher - Civil Engineering
        self.teacher_ce = User.objects.create_user(
            username='teacher_ce_user',
            email='teacher_ce@eps.edu',
            password='TeacherPass123!',
            role='teacher',
            first_name='Prof. CE'
        )
        self.teacher_ce_profile = TeacherProfile.objects.create(
            user=self.teacher_ce,
            employee_id='EMP_CE_01',
            department='Civil Engineering',
            qualification='M.Tech CE',
            institution='North Campus',
            verification_status='approved'
        )

        # Teacher - No Department Assigned
        self.teacher_nodept = User.objects.create_user(
            username='teacher_nodept_user',
            email='teacher_nodept@eps.edu',
            password='TeacherPass123!',
            role='teacher',
            first_name='Prof. Unassigned'
        )
        self.teacher_nodept_profile = TeacherProfile.objects.create(
            user=self.teacher_nodept,
            employee_id='EMP_NONE_01',
            department='',
            qualification='M.Sc',
            institution='Main Campus',
            verification_status='approved'
        )

        # Student - Pending in CS
        self.student_cs_user = User.objects.create_user(
            username='student_cs_applicant',
            email='student_cs@eps.edu',
            password='StudentPass123!',
            role='student',
            first_name='CS Applicant'
        )
        self.student_cs_profile = StudentProfile.objects.create(
            user=self.student_cs_user,
            full_name='CS Applicant',
            roll_number='ROLL_CS_01',
            department='Computer Science',
            semester='1',
            class_group=self.cls_cs,
            verification_status='pending'
        )

        # Student - Pending in CE
        self.student_ce_user = User.objects.create_user(
            username='student_ce_applicant',
            email='student_ce@eps.edu',
            password='StudentPass123!',
            role='student',
            first_name='CE Applicant'
        )
        self.student_ce_profile = StudentProfile.objects.create(
            user=self.student_ce_user,
            full_name='CE Applicant',
            roll_number='ROLL_CE_01',
            department='Civil Engineering',
            semester='1',
            class_group=self.cls_ce,
            verification_status='pending'
        )

    # ---------------------------------------------------------
    # SEC-03: Teacher Department Self-Edit Prevention Tests
    # ---------------------------------------------------------

    def test_sec03_teacher_cannot_change_department_via_profile_post(self):
        """A teacher attempting to submit a department change via POST has it strictly ignored."""
        self.client.force_login(self.teacher_cs)

        resp = self.client.post(reverse('user_profile'), {
            'full_name': 'Prof. CS Renamed',
            'email': 'cs_renamed@eps.edu',
            'department': 'Civil Engineering',  # Tampering attempt
            'qualification': 'Ph.D CS Advanced',
            'institution': 'Engineering Block A',
        })
        self.assertEqual(resp.status_code, 302)

        self.teacher_cs_profile.refresh_from_db()
        self.teacher_cs.refresh_from_db()

        # Legitimate profile fields must update
        self.assertEqual(self.teacher_cs.first_name, 'Prof. CS Renamed')
        self.assertEqual(self.teacher_cs_profile.qualification, 'Ph.D CS Advanced')
        self.assertEqual(self.teacher_cs_profile.institution, 'Engineering Block A')

        # Department must remain Computer Science (NOT Civil Engineering)
        self.assertEqual(self.teacher_cs_profile.department, 'Computer Science')

    def test_sec03_teacher_profile_form_omits_department_input(self):
        """The profile form UI displays department as read-only and has no input field."""
        self.client.force_login(self.teacher_cs)
        resp = self.client.get(reverse('user_profile'))
        self.assertEqual(resp.status_code, 200)

        # Check department is shown in read-only info, but no input field for department
        self.assertContains(resp, 'Computer Science')
        self.assertNotContains(resp, 'name="department"')

    # ---------------------------------------------------------
    # SEC-02: Cross-Department Student Approval Prevention Tests
    # ---------------------------------------------------------

    def test_sec02_teacher_cannot_approve_student_from_other_department(self):
        """Teacher from CS cannot approve a student registered in Civil Engineering."""
        self.client.force_login(self.teacher_cs)

        resp = self.client.post(
            reverse('review_student', args=[self.student_ce_profile.id]),
            {'action': 'approve', 'reason': 'Cross-dept unauthorized approve'}
        )
        # Must be rejected with HTTP 403 Forbidden
        self.assertEqual(resp.status_code, 403)

        # Student status must remain pending
        self.student_ce_profile.refresh_from_db()
        self.assertEqual(self.student_ce_profile.verification_status, 'pending')

        # No approval record should have been created
        self.assertFalse(
            ApprovalRecord.objects.filter(applicant=self.student_ce_user).exists()
        )

    def test_sec02_teacher_cannot_reject_student_from_other_department(self):
        """Teacher from CS cannot reject a student registered in Civil Engineering."""
        self.client.force_login(self.teacher_cs)

        resp = self.client.post(
            reverse('review_student', args=[self.student_ce_profile.id]),
            {'action': 'reject', 'reason': 'Cross-dept unauthorized reject'}
        )
        self.assertEqual(resp.status_code, 403)

        self.student_ce_profile.refresh_from_db()
        self.assertEqual(self.student_ce_profile.verification_status, 'pending')

    def test_sec02_teacher_cannot_view_other_department_student_review_page(self):
        """Teacher from CS cannot view student detail/review page of a CE student."""
        self.client.force_login(self.teacher_cs)

        resp = self.client.get(reverse('review_student', args=[self.student_ce_profile.id]))
        self.assertEqual(resp.status_code, 403)

    def test_sec02_teacher_approval_queue_is_strictly_scoped(self):
        """Teacher only sees applicants from their own department in the approval queue."""
        self.client.force_login(self.teacher_cs)

        resp = self.client.get(reverse('teacher_student_approvals'))
        self.assertEqual(resp.status_code, 200)

        # CS student should be present in the queryset / page
        pending_students = resp.context['pending_students']
        self.assertIn(self.student_cs_profile, pending_students)
        self.assertNotIn(self.student_ce_profile, pending_students)

    def test_sec02_teacher_with_no_department_gets_empty_queue(self):
        """Teacher without an assigned department cannot see any students or approve any."""
        self.client.force_login(self.teacher_nodept)

        resp = self.client.get(reverse('teacher_student_approvals'))
        self.assertEqual(resp.status_code, 200)

        pending_students = resp.context['pending_students']
        self.assertEqual(pending_students.count(), 0)

        # Attempt to approve CS student must fail with 403
        resp_post = self.client.post(
            reverse('review_student', args=[self.student_cs_profile.id]),
            {'action': 'approve'}
        )
        self.assertEqual(resp_post.status_code, 403)
        self.student_cs_profile.refresh_from_db()
        self.assertEqual(self.student_cs_profile.verification_status, 'pending')

    def test_sec02_legitimate_same_department_approval_succeeds(self):
        """Teacher in CS can approve pending students in CS."""
        self.client.force_login(self.teacher_cs)

        resp = self.client.post(
            reverse('review_student', args=[self.student_cs_profile.id]),
            {'action': 'approve', 'reason': 'Valid CS student verified'}
        )
        self.assertEqual(resp.status_code, 302)

        self.student_cs_profile.refresh_from_db()
        self.assertEqual(self.student_cs_profile.verification_status, 'approved')
        self.assertTrue(
            ApprovalRecord.objects.filter(applicant=self.student_cs_user, decision='approved').exists()
        )

    def test_sec02_admin_can_approve_any_department_student(self):
        """Administrator can review and approve students from any department."""
        self.client.force_login(self.admin_user)

        resp = self.client.post(
            reverse('admin_review_student', args=[self.student_ce_profile.id]),
            {'action': 'approve', 'reason': 'Admin approval for CE'}
        )
        self.assertEqual(resp.status_code, 302)

        self.student_ce_profile.refresh_from_db()
        self.assertEqual(self.student_ce_profile.verification_status, 'approved')

    # ---------------------------------------------------------
    # SEC-01 & SEC-14: Production Configuration Hardening Tests
    # ---------------------------------------------------------

    def test_sec01_production_rejects_missing_secret_key(self):
        """Production deployment strictly rejects startup if DJANGO_SECRET_KEY is missing."""
        import os, subprocess, sys
        env = os.environ.copy()
        env['DJANGO_ENV'] = 'production'
        env.pop('DJANGO_SECRET_KEY', None)

        p = subprocess.run([sys.executable, 'manage.py', 'check'], env=env, capture_output=True, text=True)
        self.assertNotEqual(p.returncode, 0)
        self.assertIn('ImproperlyConfigured', p.stderr)
        self.assertIn('DJANGO_SECRET_KEY', p.stderr)

    def test_sec01_production_rejects_wildcard_allowed_hosts(self):
        """Production deployment strictly rejects startup if DJANGO_ALLOWED_HOSTS is wildcard *."""
        import os, subprocess, sys
        env = os.environ.copy()
        env['DJANGO_ENV'] = 'production'
        env['DJANGO_SECRET_KEY'] = 'test-secret-key-that-is-at-least-fifty-characters-long-12345'
        env['DJANGO_ALLOWED_HOSTS'] = '*'

        p = subprocess.run([sys.executable, 'manage.py', 'check'], env=env, capture_output=True, text=True)
        self.assertNotEqual(p.returncode, 0)
        self.assertIn('ImproperlyConfigured', p.stderr)
        self.assertIn('wildcard', p.stderr.lower())

    def test_sec14_production_deploy_check_passes_with_safe_config(self):
        """With production settings, manage.py check --deploy passes with 0 issues."""
        import os, subprocess, sys
        env = os.environ.copy()
        env['DJANGO_ENV'] = 'production'
        env['DJANGO_SECRET_KEY'] = 'test-secret-key-that-is-at-least-fifty-characters-long-12345'
        env['DJANGO_ALLOWED_HOSTS'] = 'eps.example.edu,portal.eps.example.edu'

        p = subprocess.run([sys.executable, 'manage.py', 'check', '--deploy'], env=env, capture_output=True, text=True)
        self.assertEqual(p.returncode, 0)
        self.assertIn('System check identified no issues', p.stdout)

