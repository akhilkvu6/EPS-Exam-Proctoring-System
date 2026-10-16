from django.shortcuts import render, redirect, get_object_or_404
from django.http import JsonResponse, HttpResponse
from django.contrib import messages
from django.utils import timezone
from django.db import transaction
from django.db.models import Q, Sum
from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
from decimal import Decimal
import random
import json
import csv
import io

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import (
    Subject, Question, AnswerOption, Exam, ExamQuestion,
    ExamAttempt, StudentAnswer, ProctoringIncident
)
from .forms import SubjectForm, QuestionForm, ExamForm, RandomQuestionAddForm
from accounts.models import User, ClassGroup
from accounts.decorators import (
    teacher_required, student_required, admin_required,
    log_audit, send_notification, get_client_ip,
    check_teacher_exam_access
)


@teacher_required
def manage_exams(request):
    """List examinations created by the teacher or department."""
    teacher = request.user
    exams = Exam.objects.filter(
        Q(created_by=teacher) | Q(subject__department__name__iexact=teacher.teacher_profile.department)
    ).select_related('subject', 'created_by').prefetch_related('assigned_classes')

    now = timezone.now()
    context = {
        'exams': exams,
        'draft_count': exams.filter(status='draft').count(),
        'published_count': exams.filter(status='published').count(),
        'active_count': exams.filter(status='published', start_time__lte=now, end_time__gte=now).count(),
        'closed_count': exams.filter(Q(status='closed') | Q(status='published', end_time__lt=now)).count(),
    }
    return render(request, 'exams/manage_exams.html', context)


@teacher_required
def create_exam(request):
    """Create a new examination in Draft state."""
    teacher = request.user
    if request.method == 'POST':
        form = ExamForm(request.POST)
        if form.is_valid():
            exam = form.save(commit=False)
            exam.created_by = teacher
            exam.status = 'draft'
            exam.save()
            form.save_m2m()

            log_audit(teacher, 'EXAM_CREATED', None, request, f"Created draft exam: {exam.title}")
            messages.success(request, f"Draft examination '{exam.title}' created. Now add questions.")
            return redirect('exam_detail', exam_id=exam.id)
    else:
        # Default start and end times (tomorrow)
        initial_start = timezone.now() + timezone.timedelta(days=1)
        initial_end = initial_start + timezone.timedelta(hours=2)
        form = ExamForm(initial={
            'start_time': initial_start.strftime('%Y-%m-%dT%H:%M'),
            'end_time': initial_end.strftime('%Y-%m-%dT%H:%M'),
            'duration_minutes': 60,
            'total_marks': 100,
            'passing_marks': 40
        })

    return render(request, 'exams/create_exam.html', {'form': form})


@teacher_required
def edit_exam(request, exam_id):
    """Edit draft examination configuration."""
    exam = get_object_or_404(Exam, id=exam_id)
    if not check_teacher_exam_access(request.user, exam):
        messages.error(request, 'Access denied: You do not have authorization to edit this examination.')
        return redirect('manage_exams')

    if not exam.can_modify and not (request.user.role == 'admin' or request.user.is_superuser):
        messages.error(request, 'Published or active exams cannot be edited to maintain integrity.')
        return redirect('exam_detail', exam_id=exam.id)

    if request.method == 'POST':
        form = ExamForm(request.POST, instance=exam)
        if form.is_valid():
            form.save()
            log_audit(request.user, 'EXAM_UPDATED', None, request, f"Updated exam: {exam.title}")
            messages.success(request, f"Exam '{exam.title}' updated successfully.")
            return redirect('exam_detail', exam_id=exam.id)
    else:
        form = ExamForm(instance=exam)

    return render(request, 'exams/edit_exam.html', {'form': form, 'exam': exam})


@teacher_required
def exam_detail(request, exam_id):
    """Manage questions, settings, and review exam readiness."""
    exam = get_object_or_404(Exam, id=exam_id)
    if not check_teacher_exam_access(request.user, exam):
        messages.error(request, 'Access denied: You do not have authorization to view this examination.')
        return redirect('manage_exams')

    exam_questions = exam.exam_questions.select_related('question').all()

    # Calculate current marks total
    allocated_marks = sum(eq.effective_marks() for eq in exam_questions)

    # Questions available in question bank for this subject
    existing_q_ids = exam_questions.values_list('question_id', flat=True)
    available_questions = Question.objects.filter(subject=exam.subject).exclude(id__in=existing_q_ids)

    random_form = RandomQuestionAddForm()

    context = {
        'exam': exam,
        'exam_questions': exam_questions,
        'allocated_marks': allocated_marks,
        'available_questions': available_questions,
        'random_form': random_form,
    }
    return render(request, 'exams/exam_detail.html', context)


@teacher_required
def publish_exam(request, exam_id):
    """Publish an exam, making it visible to enrolled students."""
    exam = get_object_or_404(Exam, id=exam_id)
    if not check_teacher_exam_access(request.user, exam):
        messages.error(request, 'Access denied: You do not have authorization to publish this examination.')
        return redirect('manage_exams')

    if not exam.exam_questions.exists():
        messages.error(request, 'Cannot publish an examination with 0 questions. Please add questions first.')
        return redirect('exam_detail', exam_id=exam.id)

    exam.status = 'published'
    exam.save()
    log_audit(request.user, 'EXAM_PUBLISHED', None, request, f"Published exam: {exam.title}")

    # Notify all eligible students in-app (strictly no email)
    notified_users = set()
    for cls in exam.assigned_classes.all():
        for student_prof in cls.students.filter(verification_status='approved'):
            notified_users.add(student_prof.user)

    for student_user in exam.assigned_students.filter(student_profile__verification_status='approved'):
        notified_users.add(student_user)

    for student_user in notified_users:
        send_notification(
            recipient=student_user,
            title=f"New Examination Scheduled: {exam.title}",
            message=f"A new examination for {exam.subject.name} has been published. Schedule: {exam.start_time.strftime('%b %d, %Y %H:%M')}. Review rules and instructions in your portal.",
            notification_type="exam",
            link="/exams/my-exams/"
        )

    messages.success(request, f"Examination '{exam.title}' published successfully! {len(notified_users)} eligible students notified.")
    return redirect('exam_detail', exam_id=exam.id)


@teacher_required
def close_exam(request, exam_id):
    """
    Manually closes an examination, updating its status to 'closed'
    and safely auto-submitting all in-progress student attempts.
    """
    exam = get_object_or_404(Exam, id=exam_id)
    if not check_teacher_exam_access(request.user, exam):
        messages.error(request, 'Access denied: You do not have permission to close this examination.')
        return redirect('manage_exams')

    if exam.status == 'closed':
        messages.info(request, f"Examination '{exam.title}' is already closed.")
        return redirect('exam_detail', exam_id=exam.id)

    # Auto-submit all in-progress attempts
    active_attempts = ExamAttempt.objects.filter(exam=exam, status='in_progress')
    auto_count = 0
    for att in active_attempts:
        finalize_submission(att, request, mode='auto_submitted')
        auto_count += 1

    exam.status = 'closed'
    exam.save()

    log_audit(request.user, 'EXAM_MANUALLY_CLOSED', None, request, f"Exam '{exam.title}' closed manually. {auto_count} active attempts auto-submitted.")
    messages.success(request, f"Examination '{exam.title}' has been successfully closed. {auto_count} active attempt(s) were safely auto-submitted.")
    return redirect('exam_detail', exam_id=exam.id)


@teacher_required
def add_question_to_exam(request, exam_id):
    """Add a selected question from question bank to the exam."""
    exam = get_object_or_404(Exam, id=exam_id)
    if not check_teacher_exam_access(request.user, exam):
        messages.error(request, 'Access denied: You do not have authorization to modify this examination.')
        return redirect('manage_exams')

    if not exam.can_modify and not (request.user.role == 'admin' or request.user.is_superuser):
        messages.error(request, 'Cannot modify questions of a published exam.')
        return redirect('exam_detail', exam_id=exam.id)

    if request.method == 'POST':
        question_id = request.POST.get('question_id')
        question = get_object_or_404(Question, id=question_id)

        next_order = exam.exam_questions.count() + 1
        ExamQuestion.objects.create(
            exam=exam,
            question=question,
            order=next_order,
            marks=question.marks,
            negative_marks=question.negative_marks
        )
        messages.success(request, f"Added question to {exam.title}.")
    return redirect('exam_detail', exam_id=exam.id)


@teacher_required
def add_random_questions(request, exam_id):
    """Randomly select questions from question bank and attach to exam."""
    exam = get_object_or_404(Exam, id=exam_id)
    if not check_teacher_exam_access(request.user, exam):
        messages.error(request, 'Access denied: You do not have authorization to modify this examination.')
        return redirect('manage_exams')

    if not exam.can_modify and not (request.user.role == 'admin' or request.user.is_superuser):
        messages.error(request, 'Cannot modify questions of a published exam.')
        return redirect('exam_detail', exam_id=exam.id)

    if request.method == 'POST':
        form = RandomQuestionAddForm(request.POST)
        if form.is_valid():
            count = form.cleaned_data['count']
            difficulty = form.cleaned_data['difficulty']
            topic = form.cleaned_data['topic'].strip()

            existing_ids = exam.exam_questions.values_list('question_id', flat=True)
            qs = Question.objects.filter(subject=exam.subject).exclude(id__in=existing_ids)

            if difficulty and difficulty != 'any':
                qs = qs.filter(difficulty=difficulty)
            if topic:
                qs = qs.filter(topic__icontains=topic)

            available_pool = list(qs)
            if not available_pool:
                messages.warning(request, 'No matching unassigned questions found in the Question Bank.')
                return redirect('exam_detail', exam_id=exam.id)

            selected = random.sample(available_pool, min(count, len(available_pool)))
            current_order = exam.exam_questions.count()

            for q in selected:
                current_order += 1
                ExamQuestion.objects.create(
                    exam=exam,
                    question=q,
                    order=current_order,
                    marks=q.marks,
                    negative_marks=q.negative_marks
                )

            messages.success(request, f"Successfully selected and added {len(selected)} random questions.")
    return redirect('exam_detail', exam_id=exam.id)


@teacher_required
def remove_question_from_exam(request, exam_id, exam_question_id):
    """Remove a question from draft exam."""
    exam = get_object_or_404(Exam, id=exam_id)
    if not check_teacher_exam_access(request.user, exam):
        messages.error(request, 'Access denied: You do not have authorization to modify this examination.')
        return redirect('manage_exams')

    if not exam.can_modify and not (request.user.role == 'admin' or request.user.is_superuser):
        messages.error(request, 'Cannot modify questions of a published exam.')
        return redirect('exam_detail', exam_id=exam.id)

    eq = get_object_or_404(ExamQuestion, id=exam_question_id, exam=exam)
    eq.delete()

    # Re-order remaining questions
    for idx, remaining in enumerate(exam.exam_questions.order_by('order'), start=1):
        if remaining.order != idx:
            remaining.order = idx
            remaining.save()

    messages.info(request, 'Question removed from exam.')
    return redirect('exam_detail', exam_id=exam.id)


@teacher_required
def exam_preview(request, exam_id):
    """Preview exam as students will see it (answer keys hidden)."""
    exam = get_object_or_404(Exam, id=exam_id)
    if not check_teacher_exam_access(request.user, exam):
        messages.error(request, 'Access denied: You do not have authorization to view this examination.')
        return redirect('manage_exams')

    exam_questions = exam.exam_questions.select_related('question').prefetch_related('question__options').all()
    return render(request, 'exams/exam_preview.html', {'exam': exam, 'exam_questions': exam_questions})


@teacher_required
def question_bank(request):
    """Browse, filter, and manage reusable question repository."""
    questions = Question.objects.select_related('subject', 'created_by').prefetch_related('options').all()
    subjects = Subject.objects.all()

    # Filters
    subj_filter = request.GET.get('subject')
    type_filter = request.GET.get('type')
    diff_filter = request.GET.get('difficulty')
    search_query = request.GET.get('q', '').strip()

    if subj_filter:
        questions = questions.filter(subject_id=subj_filter)
    if type_filter:
        questions = questions.filter(question_type=type_filter)
    if diff_filter:
        questions = questions.filter(difficulty=diff_filter)
    if search_query:
        questions = questions.filter(Q(question_text__icontains=search_query) | Q(topic__icontains=search_query))

    context = {
        'questions': questions,
        'subjects': subjects,
        'selected_subject': subj_filter,
        'selected_type': type_filter,
        'selected_difficulty': diff_filter,
        'search_query': search_query,
    }
    return render(request, 'exams/question_bank.html', context)


@teacher_required
def create_question(request):
    """Create a new reusable question with options."""
    if request.method == 'POST':
        form = QuestionForm(request.POST)
        if form.is_valid():
            question = form.save(commit=False)
            question.created_by = request.user
            question.save()

            # Handle options for MCQ or True/False
            q_type = question.question_type
            if q_type == 'mcq':
                option_texts = request.POST.getlist('option_text[]')
                correct_idx = request.POST.get('correct_option')  # 0-indexed string

                for idx, opt_text in enumerate(option_texts):
                    opt_text_clean = opt_text.strip()
                    if opt_text_clean:
                        is_corr = (str(idx) == str(correct_idx))
                        AnswerOption.objects.create(
                            question=question,
                            option_text=opt_text_clean,
                            is_correct=is_corr,
                            order=idx + 1
                        )
            elif q_type == 'true_false':
                tf_correct = request.POST.get('tf_correct')  # 'True' or 'False'
                AnswerOption.objects.create(
                    question=question,
                    option_text='True',
                    is_correct=(tf_correct == 'True'),
                    order=1
                )
                AnswerOption.objects.create(
                    question=question,
                    option_text='False',
                    is_correct=(tf_correct == 'False'),
                    order=2
                )

            messages.success(request, f"Question added to Question Bank ({question.get_question_type_display()}).")
            return redirect('question_bank')
    else:
        form = QuestionForm()

    return render(request, 'exams/create_question.html', {'form': form})


@teacher_required
def edit_question(request, question_id):
    """Edit an existing question and its options with safeguards."""
    question = get_object_or_404(Question, id=question_id)
    teacher = request.user

    # Access check: admin or creator or department match
    is_admin = teacher.role == 'admin' or teacher.is_superuser or teacher.is_staff
    dept = (teacher.teacher_profile.department or '').strip().lower() if hasattr(teacher, 'teacher_profile') else ''
    subj_dept = (question.subject.department.name or '').strip().lower() if question.subject and question.subject.department else ''
    if not (is_admin or question.created_by == teacher or (dept and dept == subj_dept)):
        messages.error(request, 'Access denied: You do not have permission to edit this question.')
        return redirect('question_bank')

    has_student_answers = question.student_answers.exists()

    if request.method == 'POST':
        form = QuestionForm(request.POST, instance=question)
        if form.is_valid():
            q = form.save()

            # If question has not been answered by students, options can be reconfigured
            if not has_student_answers:
                if q.question_type == 'mcq':
                    option_texts = request.POST.getlist('option_text[]')
                    correct_idx = request.POST.get('correct_option')
                    # Remove existing options and recreate cleanly
                    q.options.all().delete()
                    for idx, opt_text in enumerate(option_texts):
                        opt_text_clean = opt_text.strip()
                        if opt_text_clean:
                            AnswerOption.objects.create(
                                question=q,
                                option_text=opt_text_clean,
                                is_correct=(str(idx) == str(correct_idx)),
                                order=idx + 1
                            )
                elif q.question_type == 'true_false':
                    tf_correct = request.POST.get('tf_correct')
                    q.options.all().delete()
                    AnswerOption.objects.create(question=q, option_text='True', is_correct=(tf_correct == 'True'), order=1)
                    AnswerOption.objects.create(question=q, option_text='False', is_correct=(tf_correct == 'False'), order=2)
            else:
                # Still allow updating text of existing options without deleting keys
                option_ids = request.POST.getlist('existing_option_id[]')
                option_texts = request.POST.getlist('existing_option_text[]')
                for opt_id, opt_txt in zip(option_ids, option_texts):
                    opt = AnswerOption.objects.filter(id=opt_id, question=q).first()
                    if opt and opt_txt.strip():
                        opt.option_text = opt_txt.strip()
                        opt.save()

            log_audit(teacher, 'QUESTION_UPDATED', None, request, f"Updated question #{q.id} ({q.subject.code})")
            messages.success(request, f"Question #{q.id} successfully updated.")
            return redirect('question_bank')
    else:
        form = QuestionForm(instance=question)

    options = question.options.order_by('order')
    return render(request, 'exams/edit_question.html', {
        'form': form,
        'question': question,
        'options': options,
        'has_student_answers': has_student_answers,
    })


@teacher_required
def delete_question(request, question_id):
    """Safeguarded question deletion from repository."""
    question = get_object_or_404(Question, id=question_id)
    teacher = request.user

    is_admin = teacher.role == 'admin' or teacher.is_superuser or teacher.is_staff
    dept = (teacher.teacher_profile.department or '').strip().lower() if hasattr(teacher, 'teacher_profile') else ''
    subj_dept = (question.subject.department.name or '').strip().lower() if question.subject and question.subject.department else ''
    if not (is_admin or question.created_by == teacher or (dept and dept == subj_dept)):
        messages.error(request, 'Access denied: You do not have permission to delete this question.')
        return redirect('question_bank')

    # Integrity safeguard: student answers check
    if question.student_answers.exists():
        messages.error(
            request,
            f"Cannot delete question #{question.id}: It is referenced by existing student attempt submissions. Question preservation is required for audit integrity."
        )
        return redirect('question_bank')

    # Integrity safeguard: check active/published exam references
    active_exam_refs = question.question_exams.filter(exam__status__in=['published', 'active'])
    if active_exam_refs.exists():
        messages.error(
            request,
            f"Cannot delete question #{question.id}: It is currently included in published examination '{active_exam_refs.first().exam.title}'."
        )
        return redirect('question_bank')

    qid = question.id
    qtext = question.question_text[:50]
    question.delete()
    log_audit(teacher, 'QUESTION_DELETED', None, request, f"Deleted question #{qid}: {qtext}")
    messages.success(request, f"Question #{qid} successfully removed from Question Bank.")
    return redirect('question_bank')


@teacher_required
def import_questions(request):
    """Bulk import questions from Excel (.xlsx) or CSV format with row validation and error reporting."""
    subjects = Subject.objects.select_related('department').all()
    import_summary = None

    if request.method == 'POST' and request.FILES.get('file'):
        uploaded_file = request.FILES['file']
        filename = uploaded_file.name.lower()
        import_mode = request.POST.get('mode', 'skip_errors')  # 'atomic' or 'skip_errors'
        default_subject_id = request.POST.get('default_subject')
        default_subject = Subject.objects.filter(id=default_subject_id).first() if default_subject_id else None

        valid_rows = []
        row_errors = []

        try:
            if filename.endswith('.xlsx'):
                wb = openpyxl.load_workbook(uploaded_file, data_only=True)
                ws = wb.active
                rows = list(ws.iter_rows(values_only=True))
                if len(rows) < 2:
                    messages.error(request, 'The uploaded Excel spreadsheet contains no data rows.')
                    return render(request, 'exams/import_questions.html', {'subjects': subjects})
                headers = [str(h).strip().lower() if h is not None else '' for h in rows[0]]
                raw_data_rows = rows[1:]
            elif filename.endswith('.csv'):
                decoded = uploaded_file.read().decode('utf-8-sig', errors='ignore')
                reader = csv.reader(io.StringIO(decoded))
                rows = list(reader)
                if len(rows) < 2:
                    messages.error(request, 'The uploaded CSV file contains no data rows.')
                    return render(request, 'exams/import_questions.html', {'subjects': subjects})
                headers = [str(h).strip().lower() for h in rows[0]]
                raw_data_rows = rows[1:]
            else:
                messages.error(request, 'Unsupported file format. Please upload an Excel (.xlsx) or CSV file.')
                return render(request, 'exams/import_questions.html', {'subjects': subjects})

            # Map column indices
            def col_idx(names):
                for name in names:
                    if name in headers:
                        return headers.index(name)
                return -1

            idx_subj = col_idx(['subject_code', 'subject', 'course_code'])
            idx_type = col_idx(['question_type', 'type'])
            idx_text = col_idx(['question_text', 'question', 'text', 'prompt'])
            idx_marks = col_idx(['marks', 'mark', 'score', 'points'])
            idx_neg = col_idx(['negative_marks', 'penalty', 'negative_mark'])
            idx_diff = col_idx(['difficulty', 'level'])
            idx_topic = col_idx(['topic', 'unit', 'module'])
            idx_opta = col_idx(['option_a', 'option1', 'opt_a'])
            idx_optb = col_idx(['option_b', 'option2', 'opt_b'])
            idx_optc = col_idx(['option_c', 'option3', 'opt_c'])
            idx_optd = col_idx(['option_d', 'option4', 'opt_d'])
            idx_corr = col_idx(['correct_option', 'correct_answer', 'answer', 'key'])
            idx_expl = col_idx(['explanation', 'solution'])

            if idx_text == -1:
                messages.error(request, "Spreadsheet missing required column: 'question_text'.")
                return render(request, 'exams/import_questions.html', {'subjects': subjects})

            for line_no, row in enumerate(raw_data_rows, start=2):
                if not any(row):
                    continue

                q_text = str(row[idx_text]).strip() if idx_text >= 0 and row[idx_text] is not None else ''
                if not q_text:
                    row_errors.append({'row': line_no, 'error': "Question text is required."})
                    continue

                # Subject resolution
                subj_code = str(row[idx_subj]).strip() if idx_subj >= 0 and row[idx_subj] is not None else ''
                subj = None
                if subj_code:
                    subj = Subject.objects.filter(Q(code__iexact=subj_code) | Q(name__iexact=subj_code)).first()
                if not subj:
                    subj = default_subject
                if not subj:
                    row_errors.append({'row': line_no, 'error': f"Subject '{subj_code}' not found and no valid default selected."})
                    continue

                # Question type
                raw_type = str(row[idx_type]).strip().lower() if idx_type >= 0 and row[idx_type] is not None else 'mcq'
                if raw_type in ['mcq', 'multiple_choice', 'multiple choice']:
                    q_type = 'mcq'
                elif raw_type in ['tf', 'true_false', 'true/false', 'true false']:
                    q_type = 'true_false'
                elif raw_type in ['short_answer', 'short', 'text']:
                    q_type = 'short_answer'
                else:
                    q_type = 'mcq'

                # Difficulty
                raw_diff = str(row[idx_diff]).strip().lower() if idx_diff >= 0 and row[idx_diff] is not None else 'medium'
                difficulty = raw_diff if raw_diff in ['easy', 'medium', 'hard'] else 'medium'

                # Marks & Negative marks
                try:
                    marks_val = float(str(row[idx_marks]).strip()) if idx_marks >= 0 and row[idx_marks] is not None else 1.0
                    marks = Decimal(str(max(0.5, marks_val)))
                except ValueError:
                    marks = Decimal('1.0')

                try:
                    neg_val = float(str(row[idx_neg]).strip()) if idx_neg >= 0 and row[idx_neg] is not None else 0.0
                    neg_marks = Decimal(str(max(0.0, neg_val)))
                except ValueError:
                    neg_marks = Decimal('0.0')

                topic = str(row[idx_topic]).strip() if idx_topic >= 0 and row[idx_topic] is not None else ''
                explanation = str(row[idx_expl]).strip() if idx_expl >= 0 and row[idx_expl] is not None else ''

                # Options validation for MCQ / True False
                opt_a = str(row[idx_opta]).strip() if idx_opta >= 0 and row[idx_opta] is not None else ''
                opt_b = str(row[idx_optb]).strip() if idx_optb >= 0 and row[idx_optb] is not None else ''
                opt_c = str(row[idx_optc]).strip() if idx_optc >= 0 and row[idx_optc] is not None else ''
                opt_d = str(row[idx_optd]).strip() if idx_optd >= 0 and row[idx_optd] is not None else ''
                corr_key = str(row[idx_corr]).strip() if idx_corr >= 0 and row[idx_corr] is not None else ''

                if q_type == 'mcq' and (not opt_a or not opt_b):
                    row_errors.append({'row': line_no, 'error': "MCQ questions require at least Option A and Option B."})
                    continue

                valid_rows.append({
                    'row_num': line_no,
                    'subject': subj,
                    'question_type': q_type,
                    'question_text': q_text,
                    'marks': marks,
                    'negative_marks': neg_marks,
                    'difficulty': difficulty,
                    'topic': topic,
                    'explanation': explanation,
                    'opt_a': opt_a,
                    'opt_b': opt_b,
                    'opt_c': opt_c,
                    'opt_d': opt_d,
                    'corr_key': corr_key,
                })

        except Exception as e:
            messages.error(request, f"Failed to parse file: {str(e)}")
            return render(request, 'exams/import_questions.html', {'subjects': subjects})

        imported_questions = []
        if import_mode == 'atomic' and row_errors:
            messages.error(request, f"Import halted in Atomic Mode. {len(row_errors)} row(s) contained errors. No questions were imported.")
        elif valid_rows:
            with transaction.atomic():
                for item in valid_rows:
                    q = Question.objects.create(
                        subject=item['subject'],
                        created_by=request.user,
                        question_type=item['question_type'],
                        difficulty=item['difficulty'],
                        topic=item['topic'],
                        question_text=item['question_text'],
                        marks=item['marks'],
                        negative_marks=item['negative_marks'],
                        explanation=item['explanation']
                    )

                    q_type = item['question_type']
                    corr = item['corr_key'].upper()

                    if q_type == 'mcq':
                        options_data = [
                            ('A', item['opt_a']),
                            ('B', item['opt_b']),
                            ('C', item['opt_c']),
                            ('D', item['opt_d']),
                        ]
                        for idx, (label, txt) in enumerate(options_data, start=1):
                            if txt:
                                is_c = (corr == label or corr == str(idx) or corr == txt.upper())
                                AnswerOption.objects.create(
                                    question=q,
                                    option_text=txt,
                                    is_correct=is_c,
                                    order=idx
                                )
                    elif q_type == 'true_false':
                        is_true_correct = corr in ['TRUE', 'T', '1', 'YES']
                        AnswerOption.objects.create(question=q, option_text='True', is_correct=is_true_correct, order=1)
                        AnswerOption.objects.create(question=q, option_text='False', is_correct=(not is_true_correct), order=2)
                    elif q_type == 'short_answer':
                        q.correct_text_answer = item['corr_key']
                        q.save()

                    imported_questions.append(q)

            log_audit(request.user, 'QUESTIONS_BULK_IMPORTED', None, request, f"Imported {len(imported_questions)} questions via bulk upload")
            send_notification(
                recipient=request.user,
                title="Questions Import Complete",
                message=f"Processed upload: {len(imported_questions)} questions added to Question Bank, {len(row_errors)} rows skipped/failed.",
                notification_type="system",
                link="/exams/questions/"
            )

            if imported_questions:
                messages.success(request, f"Successfully imported {len(imported_questions)} question(s) into the Question Bank!")
            if row_errors:
                messages.warning(request, f"{len(row_errors)} row(s) had errors and were skipped. Review the error table below.")

        import_summary = {
            'total_rows': len(valid_rows) + len(row_errors),
            'imported_count': len(imported_questions),
            'error_count': len(row_errors),
            'errors': row_errors,
            'mode': import_mode,
        }

    return render(request, 'exams/import_questions.html', {
        'subjects': subjects,
        'import_summary': import_summary,
    })


@teacher_required
def download_question_template(request):
    """Generates and downloads a standardized, styled Excel template for question bank import."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Question Bank Template"

    headers = [
        "subject_code", "question_type", "question_text", "marks", "negative_marks",
        "difficulty", "topic", "option_a", "option_b", "option_c", "option_d",
        "correct_option", "explanation"
    ]
    ws.append(headers)

    sample_rows = [
        [
            "CS101", "mcq", "Which data structure uses LIFO (Last In, First Out) ordering?",
            "2.0", "0.5", "easy", "Data Structures",
            "Queue", "Stack", "Array", "Linked List",
            "B", "A stack operates on Last-In First-Out principle."
        ],
        [
            "CS101", "true_false", "In Python, tuples are immutable data structures.",
            "1.0", "0.25", "easy", "Python Basics",
            "True", "False", "", "",
            "True", "Tuples cannot be modified after instantiation."
        ],
        [
            "CS101", "mcq", "What is the average time complexity of quicksort algorithm?",
            "2.0", "0.5", "medium", "Algorithms",
            "O(n)", "O(n log n)", "O(n^2)", "O(log n)",
            "B", "Quicksort has an average time complexity of O(n log n)."
        ],
    ]
    for row in sample_rows:
        ws.append(row)

    # Styling
    header_fill = PatternFill(start_color="2D5A27", end_color="2D5A27", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response['Content-Disposition'] = 'attachment; filename="eps-question-bank-template.xlsx"'
    return response


@teacher_required
def export_questions_csv(request):
    """Exports Question Bank repository items as a CSV file."""
    questions = Question.objects.select_related('subject', 'created_by').prefetch_related('options').all()

    subj_filter = request.GET.get('subject')
    type_filter = request.GET.get('type')
    diff_filter = request.GET.get('difficulty')
    q = request.GET.get('q', '').strip()

    if subj_filter:
        questions = questions.filter(subject_id=subj_filter)
    if type_filter:
        questions = questions.filter(question_type=type_filter)
    if diff_filter:
        questions = questions.filter(difficulty=diff_filter)
    if q:
        questions = questions.filter(Q(question_text__icontains=q) | Q(topic__icontains=q))

    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="eps-questions-export.csv"'

    writer = csv.writer(response)
    writer.writerow([
        'ID', 'Subject Code', 'Subject Name', 'Type', 'Difficulty', 'Topic',
        'Question Text', 'Marks', 'Negative Marks', 'Options Count', 'Created By', 'Created At'
    ])

    for q_obj in questions:
        writer.writerow([
            q_obj.id,
            q_obj.subject.code if q_obj.subject else '',
            q_obj.subject.name if q_obj.subject else '',
            q_obj.get_question_type_display(),
            q_obj.get_difficulty_display(),
            q_obj.topic,
            q_obj.question_text,
            q_obj.marks,
            q_obj.negative_marks,
            q_obj.options.count(),
            q_obj.created_by.username if q_obj.created_by else '',
            q_obj.created_at.strftime('%Y-%m-%d %H:%M:%S') if q_obj.created_at else ''
        ])

    return response


@teacher_required
def manage_subjects(request):
    """Create and manage academic subjects."""
    subjects = Subject.objects.select_related('department').all()
    form = SubjectForm()

    if request.method == 'POST':
        form = SubjectForm(request.POST)
        if form.is_valid():
            subj = form.save(commit=False)
            subj.created_by = request.user
            subj.save()
            messages.success(request, f"Subject '{subj.name} ({subj.code})' created.")
            return redirect('manage_subjects')

    return render(request, 'exams/manage_subjects.html', {'subjects': subjects, 'form': form})


@student_required
def student_exam_list(request):
    """Student portal view of assigned examinations."""
    student = request.user
    prof = student.student_profile
    now = timezone.now()

    # Find published exams assigned to this student or their class or department
    query = Q(status='published') & (
        Q(assigned_students=student) |
        (Q(assigned_classes=prof.class_group) if prof.class_group else Q()) |
        (Q(assigned_classes__isnull=True, assigned_students__isnull=True, subject__department__name__iexact=prof.department))
    )

    all_assigned = Exam.objects.filter(query).distinct().select_related('subject')

    active_exams = [e for e in all_assigned if e.start_time <= now <= e.end_time]
    upcoming_exams = [e for e in all_assigned if now < e.start_time]
    closed_exams = [e for e in all_assigned if now > e.end_time]

    context = {
        'active_exams': active_exams,
        'upcoming_exams': upcoming_exams,
        'closed_exams': closed_exams,
        'now': now,
    }
    return render(request, 'exams/student_exam_list.html', context)


# ==========================================
# EXAMINATION ENGINE & ATTEMPT WORKFLOW
# ==========================================

def finalize_submission(attempt, request=None, mode='submitted'):
    """Evaluates objective questions and finalizes exam attempt."""
    if attempt.status in ['submitted', 'auto_submitted']:
        return attempt

    attempt.status = mode
    attempt.submitted_at = timezone.now()

    total_score = Decimal('0.0')
    exam = attempt.exam
    exam_questions = {eq.question_id: eq for eq in exam.exam_questions.all()}

    for ans in attempt.answers.select_related('question', 'selected_option').all():
        q = ans.question
        eq = exam_questions.get(q.id)
        effective_marks = eq.effective_marks() if eq else q.marks
        effective_negative = eq.effective_negative_marks() if eq else q.negative_marks

        if q.question_type in ['mcq', 'true_false']:
            if ans.selected_option is not None:
                if ans.selected_option.is_correct:
                    ans.marks_awarded = effective_marks
                else:
                    if exam.negative_marking:
                        ans.marks_awarded = -Decimal(str(effective_negative))
                    else:
                        ans.marks_awarded = Decimal('0.0')
            else:
                ans.marks_awarded = Decimal('0.0')
            ans.is_evaluated = True
            ans.save()
            total_score += ans.marks_awarded
        else:
            ans.marks_awarded = None
            ans.is_evaluated = False
            ans.save()

    attempt.total_score = max(Decimal('0.0'), total_score)
    attempt.save()

    # Link with Result model in results app
    try:
        from results.models import Result
        result = Result.objects.filter(exam=exam, student=attempt.student).first()
        if result:
            result.attempt = attempt
            result.maximum_marks = exam.total_marks
            result.save()
        else:
            result = Result.objects.create(
                attempt=attempt,
                exam=exam,
                student=attempt.student,
                maximum_marks=exam.total_marks
            )
        result.recalculate()
    except Exception as e:
        import logging
        logging.getLogger(__name__).error(f"Error persisting result for attempt {attempt.id}: {e}")

    log_audit(
        attempt.student,
        'EXAM_SUBMITTED',
        attempt.student,
        request,
        f"Submitted exam: {exam.title} ({mode}) - Auto-score: {attempt.total_score}"
    )

    send_notification(
        recipient=attempt.student,
        title=f"Submission Recorded: {exam.title}",
        message=f"Your responses for '{exam.title}' have been safely recorded ({mode.replace('_', ' ').title()}).",
        notification_type="exam",
        link="/exams/my-exams/"
    )

    return attempt


@student_required
def exam_readiness(request, exam_id):
    """
    Candidate Examination Briefing & Pre-Flight Hardware Readiness Check.
    Displays complete exam rules, schedule, attempt allowance, negative marking scheme,
    and conducts client-side WebRTC camera/mic checks before launching the session.
    """
    exam = get_object_or_404(Exam, id=exam_id)
    student = request.user
    now = timezone.now()

    # 1. Eligibility Check
    if not exam.is_eligible(student):
        messages.error(request, 'You are not eligible to take this examination.')
        return redirect('student_exam_list')

    # Load institutional governance policies
    from accounts.policies import get_institutional_policies
    policy = get_institutional_policies()

    # Schedule status
    is_upcoming = now < exam.start_time
    is_closed = now > exam.end_time or exam.status == 'closed'
    is_window_active = not is_upcoming and not is_closed

    # Late entry check
    late_entry_enabled = policy.get('late_entry_enabled', True)
    late_entry_mins = int(policy.get('late_entry_mins', 15))
    is_late_cutoff_passed = False
    late_cutoff = None
    if late_entry_enabled and is_window_active:
        late_cutoff = exam.start_time + timezone.timedelta(minutes=late_entry_mins)
        if now > late_cutoff:
            is_late_cutoff_passed = True
    elif not late_entry_enabled and is_window_active:
        if now > exam.start_time + timezone.timedelta(seconds=60):
            is_late_cutoff_passed = True

    # Active attempt check (resumption)
    active_attempt = ExamAttempt.objects.filter(exam=exam, student=student, status='in_progress').first()

    # Completed attempts
    completed_count = ExamAttempt.objects.filter(
        exam=exam,
        student=student,
        status__in=['submitted', 'auto_submitted']
    ).count()

    policy_allowance = policy.get('attempt_allowance', 'single')
    max_policy_attempts = {'single': 1, 'dual': 2, 'practice': 3}.get(policy_allowance, 1)
    effective_attempt_limit = min(exam.attempt_limit, max_policy_attempts)
    attempts_exhausted = completed_count >= effective_attempt_limit

    # Timing accommodation calculation
    accommodation_factor = float(policy.get('timing_accommodation', '1.0'))
    adjusted_duration = int(exam.duration_minutes * accommodation_factor)

    # Questions and structure
    questions_count = exam.exam_questions.count()

    context = {
        'exam': exam,
        'active_attempt': active_attempt,
        'completed_count': completed_count,
        'effective_attempt_limit': effective_attempt_limit,
        'attempts_exhausted': attempts_exhausted,
        'is_upcoming': is_upcoming,
        'is_closed': is_closed,
        'is_window_active': is_window_active,
        'is_late_cutoff_passed': is_late_cutoff_passed,
        'late_cutoff': late_cutoff,
        'late_entry_mins': late_entry_mins,
        'adjusted_duration': adjusted_duration,
        'questions_count': questions_count,
        'policy': policy,
        'now': now,
    }
    return render(request, 'exams/exam_readiness.html', context)


@student_required
def start_exam(request, exam_id):
    """Initializes or resumes a verified student's examination attempt."""
    exam = get_object_or_404(Exam, id=exam_id)
    student = request.user
    now = timezone.now()

    # 1. Eligibility Check
    if not exam.is_eligible(student):
        messages.error(request, 'You are not eligible to take this examination.')
        return redirect('student_exam_list')

    # Load institutional governance policies
    from accounts.policies import get_institutional_policies
    policy = get_institutional_policies()

    # 2. Schedule Check & Policy Enforcement
    if now < exam.start_time:
        messages.warning(request, f"This examination opens on {exam.start_time.strftime('%b %d, %Y at %H:%M')}.")
        return redirect('student_exam_list')
    if now > exam.end_time:
        messages.error(request, 'This examination window has concluded.')
        return redirect('student_exam_list')

    # 2b. Late Entry Policy Check
    late_entry_enabled = policy.get('late_entry_enabled', True)
    late_entry_mins = int(policy.get('late_entry_mins', 15))
    if late_entry_enabled:
        late_cutoff = exam.start_time + timezone.timedelta(minutes=late_entry_mins)
        if now > late_cutoff:
            messages.error(request, f"Late entry window expired. Institutional policy permits session initiation only within {late_entry_mins} minutes of exam commencement ({exam.start_time.strftime('%H:%M')}).")
            return redirect('student_exam_list')
    else:
        if now > exam.start_time + timezone.timedelta(seconds=60):
            messages.error(request, f"Late entry is prohibited by institutional policy. This examination commenced at {exam.start_time.strftime('%H:%M')}.")
            return redirect('student_exam_list')

    # 3. Check for existing active attempt (resumption support)
    active_attempt = ExamAttempt.objects.filter(exam=exam, student=student, status='in_progress').first()
    if active_attempt:
        if now > active_attempt.server_deadline:
            finalize_submission(active_attempt, request, mode='auto_submitted')
            return redirect('exam_submitted', attempt_id=active_attempt.id)
        return redirect('exam_room', attempt_id=active_attempt.id)

    # 4. Check attempt limits with Policy Allowance Ceiling
    completed_count = ExamAttempt.objects.filter(
        exam=exam,
        student=student,
        status__in=['submitted', 'auto_submitted']
    ).count()

    policy_allowance = policy.get('attempt_allowance', 'single')
    max_policy_attempts = {'single': 1, 'dual': 2, 'practice': 3}.get(policy_allowance, 1)
    effective_attempt_limit = min(exam.attempt_limit, max_policy_attempts)

    if completed_count >= effective_attempt_limit:
        messages.error(request, f'You have reached the maximum attempt limit ({effective_attempt_limit}) for this examination.')
        return redirect('student_exam_list')

    # 5. Calculate strict server deadline with timing accommodation
    accommodation_factor = float(policy.get('timing_accommodation', '1.0'))
    adjusted_duration = int(exam.duration_minutes * accommodation_factor)
    potential_deadline = now + timezone.timedelta(minutes=adjusted_duration)
    server_deadline = min(exam.end_time, potential_deadline)

    # 6. Question sequence (with optional shuffle governed by policy)
    question_ids = list(exam.exam_questions.order_by('order').values_list('question_id', flat=True))
    if exam.shuffle_questions and policy.get('item_randomization_enabled', True):
        random.shuffle(question_ids)

    attempt = ExamAttempt.objects.create(
        exam=exam,
        student=student,
        attempt_number=completed_count + 1,
        started_at=now,
        server_deadline=server_deadline,
        ip_address=get_client_ip(request),
        user_agent=request.META.get('HTTP_USER_AGENT', '')[:250],
        shuffled_question_ids=question_ids
    )

    # Pre-populate StudentAnswer entries
    for q_id in question_ids:
        StudentAnswer.objects.get_or_create(attempt=attempt, question_id=q_id)

    log_audit(student, 'EXAM_ATTEMPT_STARTED', student, request, f"Started attempt #{attempt.attempt_number} for {exam.title}")
    return redirect('exam_room', attempt_id=attempt.id)


@student_required
def exam_room(request, attempt_id):
    """The secure, browser-proctored examination delivery interface."""
    attempt = get_object_or_404(ExamAttempt, id=attempt_id, student=request.user)
    exam = attempt.exam
    now = timezone.now()

    # If already submitted
    if attempt.status in ['submitted', 'auto_submitted']:
        return redirect('exam_submitted', attempt_id=attempt.id)

    # Server deadline enforcement
    if now > attempt.server_deadline:
        finalize_submission(attempt, request, mode='auto_submitted')
        return redirect('exam_submitted', attempt_id=attempt.id)

    from accounts.policies import get_institutional_policies
    policy = get_institutional_policies()
    distractor_shuffling = policy.get('distractor_shuffling_enabled', True)

    # Fetch questions according to attempt's deterministic sequence
    questions_dict = {q.id: q for q in Question.objects.filter(id__in=attempt.shuffled_question_ids).prefetch_related('options')}
    answers_dict = {ans.question_id: ans for ans in attempt.answers.all()}

    ordered_questions = []
    for idx, q_id in enumerate(attempt.shuffled_question_ids, start=1):
        q = questions_dict.get(q_id)
        if q:
            ans = answers_dict.get(q.id)
            options = list(q.options.all())
            if exam.shuffle_options and distractor_shuffling and q.question_type in ['mcq']:
                random.shuffle(options)

            ordered_questions.append({
                'number': idx,
                'question': q,
                'options': options,
                'answer': ans,
                'is_answered': bool(ans and (ans.selected_option_id or ans.text_answer.strip())),
                'is_marked': bool(ans and ans.is_marked_for_review),
            })

    time_remaining = attempt.time_remaining_seconds()

    context = {
        'attempt': attempt,
        'exam': exam,
        'ordered_questions': ordered_questions,
        'time_remaining': time_remaining,
        'total_questions': len(ordered_questions),
    }
    return render(request, 'exams/exam_room.html', context)


@student_required
def save_answer(request, attempt_id):
    """Autosave API endpoint for answers and marked-for-review status with ownership validation."""
    if request.method != 'POST':
        return JsonResponse({'status': 'error', 'message': 'POST required'}, status=405)

    attempt = get_object_or_404(ExamAttempt, id=attempt_id, student=request.user)

    if attempt.status != 'in_progress':
        return JsonResponse({'status': 'error', 'message': 'Exam attempt is no longer active.'}, status=403)

    # Server deadline check with 5s network grace
    now = timezone.now()
    if now > attempt.server_deadline + timezone.timedelta(seconds=5):
        finalize_submission(attempt, request, mode='auto_submitted')
        return JsonResponse({'status': 'expired', 'message': 'Exam timer has concluded.'}, status=403)

    try:
        data = json.loads(request.body.decode('utf-8'))
    except Exception:
        data = request.POST

    question_id = data.get('question_id')
    selected_option_id = data.get('selected_option_id')
    text_answer = data.get('text_answer', '')
    is_marked = data.get('is_marked_for_review')

    try:
        q_id = int(question_id)
    except (TypeError, ValueError):
        return JsonResponse({'status': 'error', 'message': 'Invalid question ID specification.'}, status=400)

    # Validate that question belongs to this attempt
    if q_id not in attempt.shuffled_question_ids:
        return JsonResponse({'status': 'error', 'message': 'Question is not part of this examination attempt.'}, status=400)

    try:
        ans, _ = StudentAnswer.objects.get_or_create(attempt=attempt, question_id=q_id)

        if selected_option_id is not None:
            if selected_option_id == '' or selected_option_id == 0 or selected_option_id == '0':
                ans.selected_option = None
            else:
                try:
                    opt_id = int(selected_option_id)
                    option = AnswerOption.objects.filter(id=opt_id, question_id=q_id).first()
                    if not option:
                        return JsonResponse({'status': 'error', 'message': 'Selected option does not belong to this question.'}, status=400)
                    ans.selected_option = option
                except (TypeError, ValueError):
                    return JsonResponse({'status': 'error', 'message': 'Invalid option ID specification.'}, status=400)

        if text_answer is not None:
            ans.text_answer = str(text_answer).strip()

        if is_marked is not None:
            ans.is_marked_for_review = bool(is_marked)

        ans.save()
        return JsonResponse({
            'status': 'saved',
            'time_remaining': attempt.time_remaining_seconds()
        })
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': str(e)}, status=400)


@student_required
def submit_exam(request, attempt_id):
    """Explicit student submission handler."""
    attempt = get_object_or_404(ExamAttempt, id=attempt_id, student=request.user)
    if request.method == 'POST':
        finalize_submission(attempt, request, mode='submitted')
        return redirect('exam_submitted', attempt_id=attempt.id)
    return redirect('exam_room', attempt_id=attempt.id)


@student_required
def exam_submitted(request, attempt_id):
    """Submission receipt and completion confirmation."""
    attempt = get_object_or_404(ExamAttempt, id=attempt_id, student=request.user)
    exam = attempt.exam

    total_q = len(attempt.shuffled_question_ids)
    answered_count = attempt.answers.filter(Q(selected_option__isnull=False) | ~Q(text_answer='')).count()
    unanswered_count = max(0, total_q - answered_count)

    context = {
        'attempt': attempt,
        'exam': exam,
        'total_questions': total_q,
        'answered_count': answered_count,
        'unanswered_count': unanswered_count,
    }
    return render(request, 'exams/exam_submitted.html', context)


@student_required
def log_proctoring_event(request, attempt_id):
    """Asynchronous telemetry receiver for client-side browser proctoring events."""
    if request.method != 'POST':
        return JsonResponse({'status': 'error'}, status=405)

    attempt = get_object_or_404(ExamAttempt, id=attempt_id, student=request.user)
    if attempt.status != 'in_progress':
        return JsonResponse({'status': 'ignored', 'reason': 'Attempt is not active'})

    try:
        data = json.loads(request.body.decode('utf-8'))
    except Exception:
        data = request.POST

    event_type = data.get('event_type')
    severity = data.get('severity', 'medium')
    details = data.get('details', '')

    incident = ProctoringIncident.objects.create(
        attempt=attempt,
        event_type=event_type,
        severity=severity,
        details=details
    )

    # Increment warning count for moderate/high violations
    if severity in ['medium', 'high']:
        attempt.warning_count += 1
        if attempt.warning_count >= attempt.exam.max_warning_count:
            attempt.proctoring_flagged = True
        attempt.save()

    return JsonResponse({
        'status': 'logged',
        'warning_count': attempt.warning_count,
        'max_warnings': attempt.exam.max_warning_count,
        'flagged': attempt.proctoring_flagged
    })


@teacher_required
def live_invigilation(request, exam_id):
    """Real-time proctoring telemetry dashboard for faculty invigilators."""
    exam = get_object_or_404(Exam, id=exam_id)
    if not check_teacher_exam_access(request.user, exam):
        messages.error(request, 'Access denied: You do not have authorization to invigilate this examination.')
        return redirect('manage_exams')

    attempts = exam.attempts.select_related('student', 'student__student_profile').prefetch_related('proctoring_incidents').all()

    now = timezone.now()
    context = {
        'exam': exam,
        'attempts': attempts,
        'in_progress_count': attempts.filter(status='in_progress').count(),
        'submitted_count': attempts.filter(status__in=['submitted', 'auto_submitted']).count(),
        'flagged_count': attempts.filter(proctoring_flagged=True).count(),
        'now': now,
    }
    return render(request, 'exams/live_invigilation.html', context)


@teacher_required
def review_incident(request, incident_id):
    """
    Review a flagged proctoring incident, record the evaluation outcome
    (Reviewed - Cleared, Inconclusive, or Requires Action), and persist review notes.
    """
    incident = get_object_or_404(ProctoringIncident, id=incident_id)
    exam = incident.attempt.exam

    if not check_teacher_exam_access(request.user, exam):
        messages.error(request, 'Access denied: You do not have permission to review incidents for this examination.')
        return redirect('manage_exams')

    if request.method == 'POST':
        status = request.POST.get('review_status', '').strip()
        notes = request.POST.get('review_notes', '').strip()

        valid_statuses = dict(ProctoringIncident.REVIEW_STATUS)
        if status in valid_statuses:
            incident.review_status = status
            incident.reviewed_by = request.user
            incident.review_notes = notes
            incident.save()

            log_audit(
                request.user,
                'INCIDENT_REVIEWED',
                incident.attempt.student,
                request,
                f"Incident #{incident.id} marked as '{valid_statuses[status]}' for {incident.attempt.student.username}"
            )
            messages.success(request, f"Incident #{incident.id} status updated to '{valid_statuses[status]}'.")
        else:
            messages.error(request, "Invalid review status outcome selected.")

        next_url = request.POST.get('next') or request.GET.get('next')
        if next_url:
            return redirect(next_url)
        return redirect('live_invigilation', exam_id=exam.id)

    return render(request, 'exams/review_incident.html', {
        'incident': incident,
        'exam': exam,
        'review_statuses': ProctoringIncident.REVIEW_STATUS,
    })


# ==============================================================================
# ADMINISTRATOR OVERSIGHT SUITE
# ==============================================================================

@admin_required
def admin_all_exams(request):
    """Institutional catalog of all examinations with filters and attempt stats."""
    from django.core.paginator import Paginator
    exams_qs = Exam.objects.all().select_related('subject', 'created_by').prefetch_related('attempts').order_by('-start_time')

    q = request.GET.get('q', '').strip()
    subject_id = request.GET.get('subject', '').strip()
    status = request.GET.get('status', '').strip()

    if q:
        exams_qs = exams_qs.filter(
            Q(title__icontains=q) |
            Q(subject__name__icontains=q) |
            Q(subject__code__icontains=q) |
            Q(created_by__username__icontains=q)
        )
    if subject_id:
        exams_qs = exams_qs.filter(subject_id=subject_id)
    if status:
        exams_qs = exams_qs.filter(status=status)

    subjects = Subject.objects.all().order_by('name')

    paginator = Paginator(exams_qs, 15)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    query_params = request.GET.copy()
    if 'page' in query_params:
        query_params.pop('page')
    preserved_params = query_params.urlencode()

    return render(request, 'exams/admin_all_exams.html', {
        'page_obj': page_obj,
        'subjects': subjects,
        'q': q,
        'subject_id': subject_id,
        'status': status,
        'preserved_params': preserved_params,
        'total_exams': exams_qs.count(),
    })


@admin_required
def admin_exam_attempts(request):
    """Institutional examination attempts oversight with candidate telemetry status."""
    from django.core.paginator import Paginator
    attempts_qs = ExamAttempt.objects.all().select_related('exam', 'exam__subject', 'student', 'student__student_profile').order_by('-started_at')

    q = request.GET.get('q', '').strip()
    status = request.GET.get('status', '').strip()
    flagged = request.GET.get('flagged', '').strip()

    if q:
        attempts_qs = attempts_qs.filter(
            Q(student__username__icontains=q) |
            Q(student__student_profile__full_name__icontains=q) |
            Q(student__student_profile__roll_number__icontains=q) |
            Q(exam__title__icontains=q)
        )
    if status:
        attempts_qs = attempts_qs.filter(status=status)
    if flagged == 'yes':
        attempts_qs = attempts_qs.filter(proctoring_flagged=True)

    paginator = Paginator(attempts_qs, 20)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    query_params = request.GET.copy()
    if 'page' in query_params:
        query_params.pop('page')
    preserved_params = query_params.urlencode()

    return render(request, 'exams/admin_attempts.html', {
        'page_obj': page_obj,
        'q': q,
        'status': status,
        'flagged': flagged,
        'preserved_params': preserved_params,
        'total_attempts': attempts_qs.count(),
    })


@admin_required
def admin_proctoring_overview(request):
    """Active proctoring monitoring overview and incident breakdown."""
    active_attempts = ExamAttempt.objects.filter(status='in_progress').select_related('exam', 'student', 'student__student_profile')
    flagged_attempts = ExamAttempt.objects.filter(proctoring_flagged=True).select_related('exam', 'student')
    recent_incidents = ProctoringIncident.objects.all().select_related('attempt', 'attempt__student', 'attempt__exam').order_by('-timestamp')[:20]

    severity_counts = {
        'high': ProctoringIncident.objects.filter(severity='high').count(),
        'medium': ProctoringIncident.objects.filter(severity='medium').count(),
        'low': ProctoringIncident.objects.filter(severity='low').count(),
        'total': ProctoringIncident.objects.count(),
    }

    return render(request, 'exams/admin_proctoring.html', {
        'active_attempts': active_attempts,
        'flagged_attempts': flagged_attempts,
        'recent_incidents': recent_incidents,
        'severity_counts': severity_counts,
    })


@admin_required
def admin_proctoring_incidents(request):
    """Filterable institutional incident ledger."""
    from django.core.paginator import Paginator
    incidents_qs = ProctoringIncident.objects.all().select_related('attempt', 'attempt__student', 'attempt__exam').order_by('-timestamp')

    event_type = request.GET.get('event_type', '').strip()
    severity = request.GET.get('severity', '').strip()
    q = request.GET.get('q', '').strip()

    if event_type:
        incidents_qs = incidents_qs.filter(event_type=event_type)
    if severity:
        incidents_qs = incidents_qs.filter(severity=severity)
    if q:
        incidents_qs = incidents_qs.filter(
            Q(attempt__student__username__icontains=q) |
            Q(attempt__exam__title__icontains=q) |
            Q(details__icontains=q)
        )

    event_types = ProctoringIncident.EVENT_TYPES

    paginator = Paginator(incidents_qs, 25)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    query_params = request.GET.copy()
    if 'page' in query_params:
        query_params.pop('page')
    preserved_params = query_params.urlencode()

    return render(request, 'exams/admin_incidents.html', {
        'page_obj': page_obj,
        'event_types': event_types,
        'event_type': event_type,
        'severity': severity,
        'q': q,
        'preserved_params': preserved_params,
        'total_incidents': incidents_qs.count(),
    })

